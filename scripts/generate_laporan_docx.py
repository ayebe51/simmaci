#!/usr/bin/env python3
"""
Generator Laporan Kerja Staff IT (Markdown to DOCX)
LP Ma'arif NU Kabupaten Cilacap
Periode: 30 Agustus - 29 September 2026
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Brand Colors (LP Ma'arif NU Emerald & Executive Palette)
COLOR_PRIMARY_HEX = "0D5C3A"       # Forest / NU Emerald
COLOR_PRIMARY = RGBColor(13, 92, 58)

COLOR_SECONDARY_HEX = "1E3A8A"     # Deep Navy
COLOR_SECONDARY = RGBColor(30, 58, 138)

COLOR_ACCENT_HEX = "D97706"        # Amber / Gold
COLOR_ACCENT = RGBColor(217, 119, 6)

COLOR_DARK_HEX = "1E293B"          # Slate 800
COLOR_DARK = RGBColor(30, 41, 59)

COLOR_MUTED_HEX = "64748B"         # Slate 500
COLOR_MUTED = RGBColor(100, 116, 139)

COLOR_LIGHT_BG_HEX = "F8FAFC"      # Slate 50
COLOR_LIGHT_BORDER_HEX = "CBD5E1"  # Slate 300
COLOR_ZEBRA_HEX = "F1F5F9"         # Slate 100
COLOR_CALLOUT_BG_HEX = "F0FDF4"    # Mint / Emerald light 50
COLOR_CALLOUT_BORDER_HEX = "10B981"# Emerald 500

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    # values in dxa (1 pt = 20 dxa)
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            val, sz, space, color = border
            borders_xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="{space}" w:color="{color}"/>'
        else:
            borders_xml += f'<w:{side} w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def format_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def format_table_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def add_inline_formatted_text(paragraph, text, default_color=COLOR_DARK, default_size=10.0, default_font="Segoe UI"):
    # Tokenize bold, italic, code, links
    pattern = re.compile(
        r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\)|🟢|🟡|🔴|✓|✗)'
    )
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 0.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+?)\]\(([^)]+?)\)', token)
            if m:
                label = m.group(1)
                run = paragraph.add_run(label)
                run.underline = True
                run.font.name = default_font
                run.font.size = Pt(default_size)
                run.font.color.rgb = COLOR_SECONDARY
            else:
                run = paragraph.add_run(token)
                run.font.name = default_font
                run.font.size = Pt(default_size)
                run.font.color.rgb = default_color
        elif token in ['🟢', '✓']:
            run = paragraph.add_run(token + " ")
            run.font.name = "Segoe UI Emoji"
            run.font.size = Pt(default_size)
            run.font.color.rgb = RGBColor(16, 185, 129)
        elif token in ['🟡']:
            run = paragraph.add_run(token + " ")
            run.font.name = "Segoe UI Emoji"
            run.font.size = Pt(default_size)
            run.font.color.rgb = COLOR_ACCENT
        elif token in ['🔴', '✗']:
            run = paragraph.add_run(token + " ")
            run.font.name = "Segoe UI Emoji"
            run.font.size = Pt(default_size)
            run.font.color.rgb = RGBColor(239, 68, 68)
        else:
            run = paragraph.add_run(token)
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color

def build_docx_from_markdown(md_path, docx_path):
    print(f"Reading markdown from: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Configure Margins (A4, 2 cm / 0.79 in)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Configure Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Laporan Kerja Staff IT — PC LP Ma'arif NU Cilacap | Periode 30 Agt – 29 Sep 2026")
    hrun.font.name = "Segoe UI"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = COLOR_MUTED

    # Configure Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Left part of footer
    frun1 = fp.add_run("Dokumen Resmi & Akuntabel • SIMMACI & Keuangan Ma'arif       ")
    frun1.font.name = "Segoe UI"
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = COLOR_MUTED

    frun2 = fp.add_run("Halaman ")
    frun2.font.name = "Segoe UI"
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = COLOR_MUTED

    fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fld1)

    frun3 = fp.add_run(" dari ")
    frun3.font.name = "Segoe UI"
    frun3.font.size = Pt(8.5)
    frun3.font.color.rgb = COLOR_MUTED

    fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    fp._p.append(fld2)

    i = 0
    total_lines = len(lines)

    in_code_block = False
    code_lang = ""
    code_lines = []

    in_table = False
    table_lines = []

    while i < total_lines:
        line = lines[i].rstrip('\r\n')

        # Check code block fences
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip().lower()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Process code block
                if code_lang == 'mermaid':
                    render_mermaid_block(doc, code_lines)
                else:
                    render_code_box(doc, code_lines, code_lang)
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Check table lines
        if line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                render_markdown_table(doc, table_lines)
                table_lines = []

        # Check empty lines
        if not line.strip():
            i += 1
            continue

        # Check Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(18)
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(13.5)
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].strip())
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(11.5)
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].strip())
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(10.5)
            run.font.color.rgb = COLOR_DARK
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="{COLOR_LIGHT_BORDER_HEX}"/>'
                f'</w:pBdr>'
            )
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Blockquote (Callout)
        if line.startswith('>'):
            callout_lines = [line[1:].strip()]
            while i + 1 < total_lines and lines[i+1].startswith('>'):
                i += 1
                callout_lines.append(lines[i][1:].strip())
            render_callout_box(doc, callout_lines)
            i += 1
            continue

        # Bullet List (* or -)
        if re.match(r'^\s*[\*\-]\s+', line):
            indent_level = len(re.match(r'^\s*', line).group(0)) // 2
            content = re.sub(r'^\s*[\*\-]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, content, default_size=10.0)
            i += 1
            continue

        # Numbered List (1., 2., etc.)
        if re.match(r'^\s*\d+\.\s+', line):
            m = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
            num = m.group(1)
            content = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            num_run = p.add_run(f"{num}. ")
            num_run.bold = True
            num_run.font.name = "Segoe UI"
            num_run.font.size = Pt(10.0)
            num_run.font.color.rgb = COLOR_PRIMARY
            add_inline_formatted_text(p, content, default_size=10.0)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_formatted_text(p, line, default_size=10.0)
        i += 1

    # In case trailing table
    if in_table and table_lines:
        render_markdown_table(doc, table_lines)

    # Save generated DOCX
    print(f"Saving DOCX to: {docx_path}")
    doc.save(docx_path)
    print("DOCX successfully generated!")

def render_code_box(doc, lines, lang):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, COLOR_ZEBRA_HEX)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    border = ('single', '4', '0', COLOR_LIGHT_BORDER_HEX)
    set_cell_borders(cell, top=border, bottom=border, left=border, right=border)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    for idx, cl in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        run = p.add_run(cl)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def render_callout_box(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, COLOR_CALLOUT_BG_HEX)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    left_border = ('single', '24', '0', COLOR_CALLOUT_BORDER_HEX)
    thin_border = ('single', '4', '0', 'E2E8F0')
    set_cell_borders(cell, top=thin_border, bottom=thin_border, left=left_border, right=thin_border)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    for idx, cl in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
        add_inline_formatted_text(p, cl, default_size=9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def render_mermaid_block(doc, lines):
    # Render structured workflow card
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0FDF4")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    border = ('single', '8', '0', COLOR_CALLOUT_BORDER_HEX)
    set_cell_borders(cell, top=border, bottom=border, left=border, right=border)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_badge = p.add_run("📊 DIAGRAM ARSITEKTUR / ALUR SISTEM (WORKFLOW):")
    run_badge.bold = True
    run_badge.font.name = "Segoe UI"
    run_badge.font.size = Pt(9.5)
    run_badge.font.color.rgb = COLOR_PRIMARY

    for cl in lines:
        if not cl.strip():
            continue
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(1)
        p2.paragraph_format.line_spacing = 1.05
        run = p2.add_run("  " + cl)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def parse_markdown_table_rows(table_lines):
    rows = []
    for l in table_lines:
        parts = [p.strip() for p in l.strip().split('|')]
        if len(parts) >= 2:
            # strip empty outer parts if line began and ended with |
            if parts[0] == '':
                parts = parts[1:]
            if parts and parts[-1] == '':
                parts = parts[:-1]
            rows.append(parts)
    return rows

def render_markdown_table(doc, table_lines):
    raw_rows = parse_markdown_table_rows(table_lines)
    if len(raw_rows) < 2:
        return

    # Check separator row
    header_row = raw_rows[0]
    data_rows = []
    for r in raw_rows[1:]:
        # is it separator row like :--- | :---:
        if all(re.match(r'^:?-+:?$', c) for c in r if c):
            continue
        data_rows.append(r)

    num_cols = len(header_row)
    if num_cols == 0:
        return

    # Normalize column count
    norm_data_rows = []
    for r in data_rows:
        if len(r) < num_cols:
            r = r + [''] * (num_cols - len(r))
        elif len(r) > num_cols:
            r = r[:num_cols]
        norm_data_rows.append(r)

    table = doc.add_table(rows=len(norm_data_rows) + 1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format Header Row
    format_table_header(table.rows[0])
    for col_idx, text in enumerate(header_row):
        cell = table.cell(0, col_idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        border = ('single', '4', '0', '0A482E')
        set_cell_borders(cell, top=border, bottom=border, left=border, right=border)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        # center header if short
        if len(text) < 15:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Format Data Rows
    for row_idx, r_data in enumerate(norm_data_rows):
        row = table.rows[row_idx + 1]
        format_table_row(row)
        is_zebra = (row_idx % 2 == 1)
        bg_color = COLOR_ZEBRA_HEX if is_zebra else "FFFFFF"

        for col_idx, text in enumerate(r_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            border = ('single', '2', '0', COLOR_LIGHT_BORDER_HEX)
            set_cell_borders(cell, top=border, bottom=border, left=border, right=border)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            # Determine alignment
            if re.match(r'^\s*(\d{1,2}/\d{1,2}/\d{4}|\d+\s*Commit|\d+\.?\d*%?|\+?\d[\d\.,]*|DONE|PASS|FIXED)\s*$', text, re.IGNORECASE):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif re.match(r'^\s*[\+\-][\d\.,]+\s*(Baris)?\s*$', text):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            add_inline_formatted_text(p, text, default_size=8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

if __name__ == '__main__':
    import sys
    md_file = sys.argv[1] if len(sys.argv) > 1 else r"d:\apss-source\SIMMACI\laporan-kerja-30-agustus-29-september-2026.md"
    docx_file = sys.argv[2] if len(sys.argv) > 2 else r"d:\apss-source\SIMMACI\laporan-kerja-30-agustus-29-september-2026.docx"
    build_docx_from_markdown(md_file, docx_file)
